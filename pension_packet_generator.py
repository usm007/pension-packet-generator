#!/usr/bin/env python3
"""Interactive retirement + Kritagyata pension document generator.

Place this script beside the two DOCX templates included in the package.
It prompts for employee details and creates:
  1. A filled four-page retirement certificate/letter packet.
  2. A filled Kritagyata data collection form.
  3. A combined DOCX containing all documents.

The script preserves the supplied blank templates and their layout. Verify all
entries against official records before signing or submission.
"""
from __future__ import annotations

import re
import sys
from calendar import monthrange
from copy import deepcopy
from datetime import date, datetime
from pathlib import Path

try:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from dateutil.relativedelta import relativedelta
except ImportError as exc:
    print("A required package is missing.")
    print("Install requirements with: python -m pip install -r requirements.txt")
    print(f"Details: {exc}")
    sys.exit(1)

HERE = Path(__file__).resolve().parent
RETIREMENT_TEMPLATE = HERE / "Retirement_Certificate_Formats_Blank_Template.docx"
KRITAGYATA_TEMPLATE = HERE / "Data_Required_for_Online_Pension_File_Kritagyata.docx"

ENCLOSURES = [
    "Single Photo : 5 copies",
    "Joint photo (husband wife) : 5 copies",
    "Family members photo 1 copy each",
    "Cancel cheque 1",
    "Bank passbook copy (Name must be same as Service Book)",
    "Non-Liability Certificate (2 copies)",
    "Service Book Verification (2 copies)",
    "Specimen Signature (2 copies)",
    "H.S.L Certificate",
    "Appointment Letter",
    "Training Passed certificate & order",
    "IPS",
    "ACPS/MACPS Order",
    "PAN Card Copy of the pensioner and family members",
    "GPF Balance sheet Copy",
]


def ask(prompt: str, default: str | None = None, required: bool = False) -> str:
    """Prompt with optional default; required answers are retried until nonblank."""
    suffix = f" [{default}]" if default not in (None, "") else ""
    while True:
        value = input(f"{prompt}{suffix}: ").strip()
        if not value and default not in (None, ""):
            value = default
        if value or not required:
            return value
        print("This field is required. Please enter a value.")


def ask_int(prompt: str, default: int, minimum: int = 0, maximum: int = 100) -> int:
    while True:
        raw = ask(prompt, str(default), required=True)
        try:
            value = int(raw)
        except ValueError:
            print("Enter a whole number.")
            continue
        if minimum <= value <= maximum:
            return value
        print(f"Enter a number from {minimum} to {maximum}.")


def parse_date(value: str, label: str, required: bool = False) -> date | None:
    formats = ("%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d")
    while True:
        value = value.strip()
        if not value and not required:
            return None
        for fmt in formats:
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                pass
        print(f"Could not read {label}. Use DD-MM-YYYY (for example, 31-12-2026).")
        value = input(f"{label}: ").strip()


def ask_date(prompt: str, default: str | None = None, required: bool = False) -> date | None:
    raw = ask(prompt, default, required=required)
    return parse_date(raw, prompt, required=required)


def fmt_date(value: date | None, slash: bool = False) -> str:
    if value is None:
        return ""
    return value.strftime("%d/%m/%Y" if slash else "%d-%m-%Y")


def formal_name(name: str) -> str:
    clean = re.sub(r"^(sri\.?\s+|shri\.?\s+)", "", name.strip(), flags=re.I)
    return f"Sri {clean}"


def slugify(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9]+", "_", value.strip()).strip("_")
    return value or "Employee"


def qualifying_service(joining: date, retirement: date) -> str:
    if retirement < joining:
        raise ValueError("The retirement date cannot be earlier than the joining date.")
    delta = relativedelta(retirement, joining)
    return f"{delta.years}Y-{delta.months}M-{delta.days}D as on {fmt_date(retirement)}"


SUPERANNUATION_AGE = 60


def superannuation_date(dob: date, age: int = SUPERANNUATION_AGE) -> date:
    """Last day of the month in which the employee attains `age` years."""
    year, month = dob.year + age, dob.month
    return date(year, month, monthrange(year, month)[1])


def age_on(dob: date, on_date: date) -> int:
    """Whole years of age on a given date."""
    return on_date.year - dob.year - ((on_date.month, on_date.day) < (dob.month, dob.day))


def period_length(from_date: date, to_date: date) -> str:
    delta = relativedelta(to_date, from_date)
    return f"{delta.years}Y-{delta.months}M-{delta.days}D"


def check_periods(periods: list[dict], joining: date, retirement: date) -> list[str]:
    """Internal consistency checks over service periods; returns warning strings."""
    warnings = []
    ordered = sorted(periods, key=lambda p: (p["from"], p["to"]))
    for i, period in enumerate(ordered):
        if period["to"] < period["from"]:
            warnings.append(f"Period {i + 1} ends before it starts.")
        if period["from"] < joining or period["to"] > retirement:
            warnings.append(f"Period {i + 1} falls outside joining–retirement dates.")
        if i and ordered[i - 1]["to"] >= period["from"]:
            warnings.append(f"Period {i} and {i + 1} overlap.")
        elif i and (period["from"] - ordered[i - 1]["to"]).days > 1:
            warnings.append(f"Gap between period {i} and {i + 1} — confirm against the Service Book.")
    return warnings


def normalize_pan(value: str) -> str:
    return re.sub(r"\s+", "", value.strip().upper())


def normalize_mobile(value: str) -> str:
    return re.sub(r"\D", "", value.strip())


def set_cell_text(cell, text: str, *, bold: bool | None = None, font_name: str | None = None,
                  font_size: float | None = None) -> None:
    """Replace cell text while keeping first paragraph/cell layout formatting."""
    p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    old_align = p.alignment
    old_style = p.style
    old_space_after = p.paragraph_format.space_after
    old_space_before = p.paragraph_format.space_before
    p.clear()
    p.style = old_style
    p.alignment = old_align
    p.paragraph_format.space_after = old_space_after
    p.paragraph_format.space_before = old_space_before
    run = p.add_run(text)
    if bold is not None:
        run.bold = bold
    if font_name:
        run.font.name = font_name
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)
    if font_size is not None:
        run.font.size = Pt(font_size)
    for extra in cell.paragraphs[1:]:
        extra.clear()


def replace_runs_in_paragraph(paragraph, replacements: dict[str, str]) -> None:
    """Replace tokens in run text, preserving run formatting where possible."""
    for run in paragraph.runs:
        text = run.text
        for old, new in replacements.items():
            if old in text:
                text = text.replace(old, new)
        # Assign only when text changes; touching an empty break-only run can
        # inadvertently delete a manual page break in the supplied template.
        if text != run.text:
            run.text = text


def replace_in_document(doc, replacements: dict[str, str]) -> None:
    for paragraph in doc.paragraphs:
        replace_runs_in_paragraph(paragraph, replacements)
    for table in doc.tables:
        for row in table.rows:
            # Merged cells can appear more than once, so duplicate work is harmless.
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_runs_in_paragraph(paragraph, replacements)


def clone_table_row(table, source_row_index: int = -1):
    """Append a copy of an existing row, preserving formatting and cell properties."""
    tr = deepcopy(table.rows[source_row_index]._tr)
    table._tbl.append(tr)
    return table.rows[-1]


def fill_retirement_template(data: dict, service_periods: list[dict], out_path: Path) -> None:
    if not RETIREMENT_TEMPLATE.exists():
        raise FileNotFoundError(f"Template not found: {RETIREMENT_TEMPLATE}")
    doc = Document(RETIREMENT_TEMPLATE)
    full_name = formal_name(data["name"])
    display_dates = {
        "[Name of Employee]": full_name,
        "[Designation]": data["designation"],
        "[GPF Account No.]": data["gpf"],
        "[Employee Code No.]": data["employee_code"],
        "[Date of Joining]": fmt_date(data["joining"]),
        "[Date of Superannuation]": fmt_date(data["retirement"]),
        "[Date of Issue]": fmt_date(data["issue_date"], slash=True),
    }
    replace_in_document(doc, display_dates)

    # Service Verification Certificate's employee particulars table.
    employee_table = doc.tables[0]
    values = [
        full_name,
        data["designation"],
        data["employee_code"],
        fmt_date(data["dob"]),
        fmt_date(data["joining"]),
        fmt_date(data["retirement"]),
        data["qualifying_service"],
    ]
    for row, value in zip(employee_table.rows, values):
        set_cell_text(row.cells[1], ":  " + value)

    # Service-book verification periods: retain the template's 11 rows, add more when needed.
    verification_table = doc.tables[1]
    desired_data_rows = max(11, len(service_periods))
    while len(verification_table.rows) < desired_data_rows + 2:
        clone_table_row(verification_table, -1)
    for i in range(desired_data_rows):
        row = verification_table.rows[i + 2]
        if i < len(service_periods):
            period = service_periods[i]
            from_text = fmt_date(period["from"], slash=True)
            to_text = fmt_date(period["to"], slash=True)
            page_text = period["page"]
        else:
            from_text = to_text = page_text = ""
        set_cell_text(row.cells[0], str(i + 1))
        set_cell_text(row.cells[1], from_text)
        set_cell_text(row.cells[2], to_text)
        set_cell_text(row.cells[3], page_text)

    # Specimen signature table: names in capitals, with signature spaces left blank.
    signature_table = doc.tables[2]
    upper_name = re.sub(r"^(SRI|SHRI)\s+", "", full_name, flags=re.I).upper()
    for i in range(1, len(signature_table.rows)):
        set_cell_text(signature_table.rows[i].cells[1], upper_name, bold=True)
    # Replace the attesting officer's placeholder name only within the merged attestation cell.
    attestation_cell = signature_table.rows[1].cells[3]
    if data["headmaster"]:
        for paragraph in attestation_cell.paragraphs:
            for run in paragraph.runs:
                if "______________" in run.text:
                    run.text = run.text.replace("______________", data["headmaster"])

    # Signature blocks and letter signature: add the printed Headmaster/DDO name if supplied.
    for paragraph_index in (13, 25, 67):
        if paragraph_index < len(doc.paragraphs) and data["headmaster"]:
            p = doc.paragraphs[paragraph_index]
            for run in p.runs:
                if "______________" in run.text:
                    run.text = run.text.replace("______________", data["headmaster"])
                    break

    doc.core_properties.title = f"Retirement Documents - {data['name']}"
    doc.core_properties.subject = "Retirement certificates and letter generated from a reusable template"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(out_path)


def fill_kritagyata_template(data: dict, family: list[dict], enclosed_numbers: set[int], out_path: Path) -> None:
    if not KRITAGYATA_TEMPLATE.exists():
        raise FileNotFoundError(f"Template not found: {KRITAGYATA_TEMPLATE}")
    doc = Document(KRITAGYATA_TEMPLATE)
    # Main employee information: third cell of each row is the entry space.
    employee_table = doc.tables[0]
    values = [
        data["pan"],
        formal_name(data["name"]),
        data["mobile"],
        data["email"],
        fmt_date(data["dob"]),
        fmt_date(data["joining"]),
        fmt_date(data["retirement"]),
        data["gpf"],
        data["height"],
        data["blood_group"],
        data["address"],
        data["identification_mark"],
        data["treasury"],
    ]
    for row, value in zip(employee_table.rows, values):
        set_cell_text(row.cells[2], value, font_name="Arial", font_size=7.5)

    # Fill family-member rows; add rows if family contains more than the original three.
    family_table = doc.tables[1]
    while len(family_table.rows) < len(family) + 1:
        clone_table_row(family_table, -1)
    for i in range(max(3, len(family))):
        row = family_table.rows[i + 1]
        if i < len(family):
            member = family[i]
            values = [
                str(i + 1), member["name"], member["dob"], member["marital_status"],
                member["relation"], member["disability"], member["blood_group"],
                member["height"], member["id_proof"],
            ]
        else:
            values = [str(i + 1), "", "", "", "", "", "", "", ""]
        for j, value in enumerate(values):
            set_cell_text(row.cells[j], value, font_name="Arial", font_size=6.8 if j > 0 else 7)

    enclosure_table = doc.tables[2]
    for i, row in enumerate(enclosure_table.rows, start=1):
        set_cell_text(row.cells[1], "Enclosed" if i in enclosed_numbers else "Pending",
                      font_name="Arial", font_size=7.5)

    doc.core_properties.title = f"Kritagyata Pension File Data - {data['name']}"
    doc.core_properties.subject = "Online pension-file data form generated from a reusable template"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(out_path)


def combine_documents(retirement_path: Path, kritagyata_path: Path, out_path: Path) -> None:
    """Combine the retirement packet and Kritagyata form, preserving a new page section."""
    combined = Document(retirement_path)
    extra = Document(kritagyata_path)
    source_section = extra.sections[-1]
    # Start the Kritagyata form on its own page with the original form's wider content area.
    new_section = combined.add_section(WD_SECTION.NEW_PAGE)
    new_section.page_width = source_section.page_width
    new_section.page_height = source_section.page_height
    new_section.top_margin = source_section.top_margin
    new_section.bottom_margin = source_section.bottom_margin
    new_section.left_margin = source_section.left_margin
    new_section.right_margin = source_section.right_margin
    new_section.header_distance = source_section.header_distance
    new_section.footer_distance = source_section.footer_distance

    target_body = combined.element.body
    final_sect_pr = target_body.sectPr
    for child in extra.element.body:
        if child.tag == qn("w:sectPr"):
            continue
        target_body.insert(target_body.index(final_sect_pr), deepcopy(child))

    combined.core_properties.title = "Complete Retirement and Kritagyata Pension Packet"
    combined.core_properties.subject = "Combined retirement certificates, letter and Kritagyata data form"
    combined.core_properties.author = ""
    combined.core_properties.last_modified_by = ""
    combined.save(out_path)


def gather_information() -> tuple[dict, list[dict], list[dict], set[int]]:
    print("\nRETIREMENT + KRITAGYATA PENSION PACKET GENERATOR")
    print("Enter employee details once. Press Enter for optional information you do not have.")
    print("Dates can be entered as DD-MM-YYYY or DD/MM/YYYY.\n")

    name = ask("Employee full name (without Sri/Shri)", required=True)
    designation = ask("Designation", required=True)
    employee_code = ask("Employee Code No. (optional)")
    gpf = ask("GPF Account No.")
    dob = ask_date("Date of Birth (DD-MM-YYYY)", required=True)
    joining = ask_date("Date of Joining (DD-MM-YYYY)", required=True)
    suggested = fmt_date(superannuation_date(dob)) if dob else None
    if suggested:
        print(f"Suggested superannuation ({SUPERANNUATION_AGE} years, last day of month): {suggested} — Enter accepts it.")
    retirement = ask_date("Date of Superannuation/Retirement (DD-MM-YYYY)", default=suggested, required=True)
    if retirement < joining:
        raise ValueError("Retirement date cannot be earlier than the joining date.")
    print(f"Age at retirement: {age_on(dob, retirement)} years.")
    issue_default = date.today().strftime("%d/%m/%Y")
    issue_date = ask_date("Date of Issue", default=issue_default, required=True)
    headmaster = ask("Headmaster/DDO printed name (optional)")

    calculated = qualifying_service(joining, retirement)
    print(f"\nCalculated qualifying service from the dates entered: {calculated}")
    print("Confirm this calculation against the official Service Book before issuing.")
    qualifying = ask("Qualifying service override (press Enter to use calculated value)") or calculated

    period_count = ask_int("Number of service verification periods to enter", 0, 0, 100)
    periods = []
    for i in range(period_count):
        print(f"\nService verification period {i + 1}")
        from_date = ask_date("  From date (DD/MM/YYYY)", required=True)
        to_date = ask_date("  To date (DD/MM/YYYY)", required=True)
        page = ask("  Page No. in Service Book", required=True)
        periods.append({"from": from_date, "to": to_date, "page": page})
    for i, period in enumerate(periods, start=1):
        print(f"  Period {i} length: {period_length(period['from'], period['to'])}")
    for warning in check_periods(periods, joining, retirement):
        print(f"  Note: {warning}")

    print("\nKRITAGYATA ONLINE PENSION FILE DETAILS")
    pan = normalize_pan(ask("PAN No."))
    mobile = normalize_mobile(ask("Mobile No."))
    email = ask("Email ID")
    height = ask("Employee height in cm (as per Service Book)")
    blood_group = ask("Employee blood group")
    address = ask("Present full address")
    identification_mark = ask("Personal identification mark (as per Service Book)")
    treasury = ask("Name of Treasury")

    family_count = ask_int("Number of family members to enter in the form", 0, 0, 100)
    family = []
    for i in range(family_count):
        print(f"\nFamily member {i + 1}")
        member = {
            "name": ask("  Full name", required=True),
            "dob": ask("  Date of birth as per PAN (DD-MM-YYYY)"),
            "marital_status": ask("  Marital status"),
            "relation": ask("  Relation to pensioner"),
            "disability": ask("  Disability (Yes/No)", default="No"),
            "blood_group": ask("  Blood group"),
            "height": ask("  Height in centimetres"),
            "id_proof": ask("  ID proof enclosed (e.g. PAN Card / HSLC Certificate / Birth Certificate)"),
        }
        if member["dob"]:
            parsed_member_dob = parse_date(member["dob"], "Family member DOB")
            member["dob"] = fmt_date(parsed_member_dob)
        family.append(member)

    print("\nEnclosure checklist (select the numbers of documents already enclosed):")
    for i, item in enumerate(ENCLOSURES, start=1):
        print(f"  {i:>2}. {item}")
    while True:
        raw = ask("Enter enclosed item numbers separated by commas (Enter = none)")
        try:
            enclosed = {int(piece.strip()) for piece in raw.split(",") if piece.strip()}
        except ValueError:
            print("Use numbers separated by commas, for example: 1,4,6,8")
            continue
        if all(1 <= number <= len(ENCLOSURES) for number in enclosed):
            break
        print(f"Numbers must be between 1 and {len(ENCLOSURES)}.")

    data = {
        "name": name, "designation": designation, "employee_code": employee_code,
        "gpf": gpf, "dob": dob, "joining": joining, "retirement": retirement,
        "issue_date": issue_date, "headmaster": headmaster,
        "qualifying_service": qualifying, "pan": pan, "mobile": mobile, "email": email,
        "height": height, "blood_group": blood_group, "address": address,
        "identification_mark": identification_mark, "treasury": treasury,
    }
    return data, periods, family, enclosed


def main() -> int:
    if not RETIREMENT_TEMPLATE.exists() or not KRITAGYATA_TEMPLATE.exists():
        print("Could not find the DOCX templates beside the script.")
        print("Keep these files in the same folder as pension_packet_generator.py:")
        print(f"  - {RETIREMENT_TEMPLATE.name}")
        print(f"  - {KRITAGYATA_TEMPLATE.name}")
        return 2

    try:
        data, periods, family, enclosed = gather_information()
        output_dir_raw = ask("Output folder", default="Generated_Pension_Packets")
        output_dir = Path(output_dir_raw).expanduser()
        if not output_dir.is_absolute():
            output_dir = HERE / output_dir
        output_dir.mkdir(parents=True, exist_ok=True)

        safe_name = slugify(data["name"])
        retirement_out = output_dir / f"Retirement_Certificates_{safe_name}.docx"
        kritagyata_out = output_dir / f"Kritagyata_Form_{safe_name}.docx"
        combined_out = output_dir / f"Complete_Pension_Packet_{safe_name}.docx"

        fill_retirement_template(data, periods, retirement_out)
        fill_kritagyata_template(data, family, enclosed, kritagyata_out)
        combine_documents(retirement_out, kritagyata_out, combined_out)

        print("\nCompleted successfully. Files created:")
        print(f"  Retirement certificates + letter: {retirement_out}")
        print(f"  Kritagyata form:                 {kritagyata_out}")
        print(f"  Combined packet:                 {combined_out}")
        print("\nReminder: verify all details, qualifying service and enclosure status against official records before signing/submission.")
        return 0
    except (KeyboardInterrupt, EOFError):
        print("\nCancelled. No further files were generated.")
        return 130
    except Exception as exc:
        print(f"\nError: {exc}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
